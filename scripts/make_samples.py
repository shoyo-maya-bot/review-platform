"""samples/ の架空の設計書を作り直す（python scripts/make_samples.py）. 中身は架空の会員 API で、わざと誤りを 3 つ入れてある."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

OUT = Path(__file__).resolve().parents[1] / "samples"

ROWS = [
    ("No", "項目名", "物理名", "型", "桁", "必須", "備考"),
    (1, "会員ID", "member_id", "数値", 10, "○", "採番"),
    (2, "氏名", "name", "文字列", 50, "○", "前後の空白を除く"),
    (3, "メールアドレス", "email", "文字列", 254, "○", "重複不可（大文字小文字を区別しない）"),
    (4, "電話番号", "tel", "文字列", 11, "", "ハイフンなし"),
    (5, "生年月日", "birth_date", "日付", 8, "○", "YYYYMMDD"),
    (6, "会員区分", "member_type", "コード", 1, "○", "1: 一般 / 2: 法人 / 9: その他（など）"),
]


def main() -> None:
    OUT.mkdir(exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "概要"
    ws["A1"] = "API 設計書: 会員登録（API0002）"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A3"], ws["B3"] = "メソッド", "POST /api/v1/members"
    ws["A4"], ws["B4"] = "正常", "201 Created。Location ヘッダーに照会の URL"
    ws["A5"], ws["B5"] = "入力の誤り", "400。メッセージ ID APPE0001"
    ws["A6"], ws["B6"] = "重複", "409。メッセージ ID APPE0003"
    # わざと入れた誤り 1: 電話番号の桁が項目定義（11）と違う
    ws["A7"], ws["B7"] = "電話番号", "国際番号を含めて最大 15 桁"
    items = wb.create_sheet("項目定義")
    for r in ROWS:
        items.append(r)
    for c in items[1]:
        c.font = Font(bold=True)
    msgs = wb.create_sheet("メッセージ")
    msgs.append(("ID", "本文"))
    msgs.append(("APPE0001", "入力内容に誤りがあります。"))
    # わざと入れた誤り 2: 概要で使う APPE0003 がメッセージの表に無い
    msgs.append(("APPE0002", "会員が見つかりません。"))
    wb.save(OUT / "member-api-design.xlsx")

    md = ["# API 設計書: 会員登録（API0002）", "", "## 項目定義", "", "| " + " | ".join(map(str, ROWS[0])) + " |", "|" + "---|" * len(ROWS[0])]
    md += ["| " + " | ".join(map(str, r)) + " |" for r in ROWS[1:]]
    # わざと入れた誤り 3: 生年月日の未来日・存在しない日付の扱いが書かれていない
    md += ["", "## 処理", "", "1. 入力を検証する（誤りは 400）", "2. メールアドレスの重複を確かめる（重複は 409）", "3. 登録して 201 を返す", ""]
    (OUT / "member-api-design.md").write_text("\n".join(md), encoding="utf-8")


if __name__ == "__main__":
    main()
